"""Aperçu des documents bureautiques (ERP, extranet, espace client) : tableurs, Word, PowerPoint, OpenDocument, CSV
— conversion sans stockage, HTML assaini, archives contrôlées, accès réservé aux utilisateurs connectés."""

from __future__ import annotations

import io
import zipfile
from typing import Any

from tests.conftest import as_role

API = "/api/v1"


def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Métré"
    ws.append(["Désignation", "Unité", "Quantité", "PU"])
    ws.append(["Béton dosé à 350 kg", "m3", 42.5, 95000])
    ws.append(["Acier HA 12", "kg", 1800, 850])
    ws.append(["Total", "", "", "=SUM(D2:D3)"])
    wb.create_sheet("Planning").append(["Lot", "Début"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Compte rendu de chantier n°12", level=1)
    paragraph = document.add_paragraph("Coulage du plancher ")
    paragraph.add_run("haut R+1").bold = True
    document.add_paragraph("Contrôle des armatures", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Lot"
    table.cell(0, 1).text = "Avancement"
    table.cell(1, 0).text = "Gros œuvre <script>alert(1)</script>"
    table.cell(1, 1).text = "65 %"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _zip(files: dict[str, str]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return buf.getvalue()


def _pptx() -> bytes:
    slide = (
        '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:cSld><p:spTree>'
        "<p:sp><p:txBody><a:p><a:r><a:t>{title}</a:t></a:r></a:p><a:p><a:r><a:t>{line}</a:t></a:r></a:p>"
        "</p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
    )
    return _zip(
        {
            "ppt/slides/slide2.xml": slide.format(title="Planning", line="Livraison en mars"),
            "ppt/slides/slide1.xml": slide.format(title="Présentation du projet", line="Villa R+1"),
            "ppt/slides/slide10.xml": slide.format(title="Annexes", line="Plans"),
        }
    )


def _ods() -> bytes:
    content = (
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:spreadsheet>'
        '<table:table table:name="Budget"><table:table-row><table:table-cell><text:p>Poste</text:p></table:table-cell>'
        "<table:table-cell><text:p>Montant</text:p></table:table-cell></table:table-row>"
        '<table:table-row table:number-rows-repeated="2"><table:table-cell><text:p>Terrassement</text:p>'
        '</table:table-cell><table:table-cell table:number-columns-repeated="1"><text:p>1 200 000</text:p>'
        "</table:table-cell></table:table-row></table:table></office:spreadsheet></office:body></office:document-content>"
    )
    return _zip({"content.xml": content, "mimetype": "application/vnd.oasis.opendocument.spreadsheet"})


def _upload(client: Any, headers: dict[str, str], name: str, data: bytes) -> Any:
    return client.post(f"{API}/files/preview", headers=headers, files={"file": (name, data)})


def test_spreadsheets_and_csv(client: Any, db: Any) -> None:
    _user, headers = as_role(client, db, "CHEF_PROJET")
    body = _upload(client, headers, "metre.xlsx", _xlsx()).json()
    assert body["kind"] == "table" and [s["name"] for s in body["sheets"]] == ["Métré", "Planning"]
    first = body["sheets"][0]
    assert first["rows"][1] == ["Béton dosé à 350 kg", "m3", "42.5", "95000"] and first["total_rows"] == 4
    assert first["rows"][3][3] == "ƒ =SUM(D2:D3)", "formule affichée si le classeur n'a jamais été recalculé"

    csv_bytes = "Lot;Montant\r\nGros œuvre;12 500 000\r\nToiture;4 200 000\r\n".encode("cp1252")
    body = _upload(client, headers, "budget.csv", csv_bytes).json()
    assert body["delimiter"] == ";" and body["sheets"][0]["rows"][1] == ["Gros œuvre", "12 500 000"]

    body = _upload(client, headers, "budget.ods", _ods()).json()
    assert body["sheets"][0]["name"] == "Budget" and body["sheets"][0]["total_rows"] == 3


def test_word_and_slides_are_sanitized(client: Any, db: Any) -> None:
    _user, headers = as_role(client, db, "CLIENT")
    body = _upload(client, headers, "compte-rendu.docx", _docx()).json()
    assert body["kind"] == "html"
    assert "<strong>haut R+1</strong>" in body["html"] and "<li>Contrôle des armatures</li>" in body["html"]
    assert "<script>" not in body["html"] and "&lt;script&gt;" in body["html"]

    body = _upload(client, headers, "presentation.pptx", _pptx()).json()
    assert body["kind"] == "slides" and [s["title"] for s in body["slides"]] == [
        "Présentation du projet",
        "Planning",
        "Annexes",
    ]


def test_guards(client: Any, db: Any, monkeypatch: Any) -> None:
    from app.core import preview as preview_module

    anonymous = _upload(client, {}, "metre.xlsx", _xlsx())
    assert anonymous.status_code == 401
    _user, headers = as_role(client, db, "FOURNISSEUR")
    unsupported = _upload(client, headers, "plan.dwg", b"AC1032" + b"\0" * 64)
    assert unsupported.status_code == 422 and unsupported.json()["error"]["code"] == "PREVIEW_UNSUPPORTED"
    broken = _upload(client, headers, "rapport.docx", b"PK\x03\x04 pas un vrai docx")
    assert broken.status_code == 422 and broken.json()["error"]["code"] == "PREVIEW_UNAVAILABLE"
    monkeypatch.setattr(preview_module, "MAX_UNZIPPED_BYTES", 4096)
    bomb = _zip({"content.xml": "0" * 10_000})
    assert _upload(client, headers, "bombe.ods", bomb).status_code == 422
