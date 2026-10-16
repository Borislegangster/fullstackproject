"""GED (§12) : aucune version écrasée, cycle brouillon → revue → correction → validé → publié → obsolète →
archivé ; un plan Vn devient obsolète dès la publication de Vn+1 (V4-GED-29) ; verrouillage pendant
approbation ; signature OTP liée à l'empreinte ; filigrane ; QR d'authenticité ; recherche plein texte ;
comparaison texte ; suppression logique avec justification ; journal d'accès (V4-GED-30)."""

from __future__ import annotations

import difflib
import io
import uuid
from typing import Any

from sqlalchemy import func, or_, select, text
from sqlalchemy.orm import Session

from app.core import audit
from app.core.access import Access
from app.core.errors import BusinessRuleError, Conflict, Forbidden, NotFound
from app.core.events import emit
from app.core.files import ValidatedUpload, store_private
from app.core.numbering import next_number
from app.core.security import random_token
from app.core.utils import utcnow
from app.modules.documents.models import (
    Document,
    DocumentAccessLog,
    DocumentApprovalStep,
    DocumentLink,
    DocumentPermission,
    DocumentVersion,
    Signature,
)

VERSION_FLOW = {
    "draft": {"in_review", "validated", "published", "archived"},
    "in_review": {"correction", "validated"},
    "correction": {"in_review", "draft"},
    "validated": {"published", "correction"},
    "published": {"obsolete", "archived"},
    "obsolete": {"archived"},
}


def extract_text(upload: ValidatedUpload) -> tuple[str | None, int | None]:
    from app.core.pdf import extract_text as pdf_text

    if upload.mime == "application/pdf":
        text_value, pages = pdf_text(upload.data)
        return text_value or None, pages
    if upload.mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        try:
            import docx

            document = docx.Document(io.BytesIO(upload.data))
            return "\n".join(p.text for p in document.paragraphs)[:500_000], None
        except Exception:  # noqa: BLE001
            return None, None
    if upload.mime in ("text/plain", "text/csv"):
        return upload.data.decode("utf-8", errors="replace")[:500_000], None
    if upload.mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        try:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(upload.data), read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets[:5]:
                for row in ws.iter_rows(max_row=2000, values_only=True):
                    rows.append(" ".join(str(c) for c in row if c is not None))
            return "\n".join(rows)[:500_000], None
        except Exception:  # noqa: BLE001
            return None, None
    return None, None


# ---------------------------------------------------------------------------- contrôle d'accès
def can_read(access: Access, db: Session, doc: Document) -> bool:
    if doc.deleted_at is not None:
        return False
    acl = (
        db.execute(
            select(DocumentPermission).where(
                or_(DocumentPermission.document_id == doc.id, DocumentPermission.folder_id == doc.folder_id)
            )
        )
        .scalars()
        .all()
    )
    deny = any(p.permission == "deny" and _matches(p, access) for p in acl)
    if deny:
        return False
    explicit = any(p.permission in ("read", "write", "approve", "sign") and _matches(p, access) for p in acl)
    if access.p.is_client:
        return (
            bool(doc.client_visible and doc.published_version_id and access.can_access_project(doc.project_id))
            or explicit
        )
    if access.p.org_kind in ("supplier", "subcontractor", "partner"):
        return (
            bool(doc.partner_visible and doc.published_version_id and access.can_access_project(doc.project_id))
            or explicit
        )
    if not access.has("documents:read"):
        return explicit
    if doc.project_id and not access.can_access_project(doc.project_id):
        return explicit
    if not access.can_read_confidentiality(doc.confidentiality):
        return explicit and doc.confidentiality != "very_sensitive"
    disciplines = access.allowed_disciplines(doc.project_id) if doc.project_id else None
    if disciplines is not None and doc.discipline and doc.discipline not in disciplines:
        return explicit
    lots = access.allowed_lots(doc.project_id) if doc.project_id else None
    return not (lots is not None and doc.lot_id and doc.lot_id not in lots) or explicit


def _matches(perm: DocumentPermission, access: Access) -> bool:
    value = perm.principal_value
    if perm.principal_type == "user":
        return value == str(access.p.id)
    if perm.principal_type == "role":
        return value in access.p.roles
    if perm.principal_type == "organization":
        return value == str(access.p.organization_id)
    return False


def require_read(access: Access, db: Session, doc: Document) -> None:
    if not can_read(access, db, doc):
        raise NotFound("Document introuvable")


def require_write(access: Access, db: Session, doc: Document) -> None:
    require_read(access, db, doc)
    access.require("documents:write")
    if doc.project_id:
        access.require_project(doc.project_id)


# ---------------------------------------------------------------------------- création & versions
def create_document(
    db: Session,
    *,
    access: Access,
    upload: ValidatedUpload,
    title: str,
    doc_type: str,
    project_id: uuid.UUID | None,
    site_id: uuid.UUID | None,
    lot_id: uuid.UUID | None,
    folder_id: uuid.UUID | None,
    discipline: str | None,
    confidentiality: str,
    description: str | None,
    tags: list[str],
    code: str | None = None,
) -> tuple[Document, DocumentVersion]:
    if project_id:
        access.require_project(project_id, "documents:write")
    else:
        access.require("documents:write")
    if not access.can_read_confidentiality(confidentiality):
        raise Forbidden("Niveau de confidentialité non autorisé pour votre profil")
    document = Document(
        code=code or next_number(db, "document"),
        title=title,
        description=description,
        doc_type=doc_type,
        discipline=discipline,
        project_id=project_id,
        site_id=site_id,
        lot_id=lot_id,
        folder_id=folder_id,
        confidentiality=confidentiality,
        tags=tags,
        created_by_id=access.p.id,
    )
    db.add(document)
    db.flush()
    version = _add_version(db, document, upload, author_id=access.p.id, major=1, minor=0, parent=None, reason=None)
    audit.record(
        db,
        family="business",
        module="documents",
        action="CREATE",
        resource_type="document",
        resource_id=document.id,
        project_id=project_id,
        diff={
            "code": document.code,
            "version": version.label,
            "checksum": version.checksum_sha256,
            "confidentiality": confidentiality,
        },
    )
    return document, version


def _add_version(
    db: Session,
    document: Document,
    upload: ValidatedUpload,
    *,
    author_id: uuid.UUID,
    major: int,
    minor: int,
    parent: DocumentVersion | None,
    reason: str | None,
    source: str = "upload",
) -> DocumentVersion:
    stored = store_private(
        db, upload, purpose="documents", classification=document.confidentiality, uploaded_by_id=author_id
    )
    text_value, pages = extract_text(upload)
    version = DocumentVersion(
        document_id=document.id,
        version_major=major,
        version_minor=minor,
        parent_version_id=parent.id if parent else None,
        file_id=stored.id,
        checksum_sha256=upload.sha256,
        size_bytes=upload.size,
        mime_type=upload.mime,
        original_name=upload.filename,
        author_id=author_id,
        source=source,
        status="draft",
        revision_reason=reason,
        extracted_text=text_value,
        page_count=pages,
        authenticity_token=random_token(24),
    )
    db.add(version)
    db.flush()
    document.current_version_id = version.id
    return version


def new_version(
    db: Session, *, access: Access, document: Document, upload: ValidatedUpload, reason: str, major: bool
) -> DocumentVersion:
    require_write(access, db, document)
    if document.locked:
        raise Conflict("Document verrouillé (approbation en cours)", code="DOCUMENT_LOCKED")
    if not reason or len(reason.strip()) < 3:
        raise BusinessRuleError("Le motif de révision est obligatoire", code="REVISION_REASON_REQUIRED")
    current = db.get(DocumentVersion, document.current_version_id) if document.current_version_id else None
    if current and current.checksum_sha256 == upload.sha256:
        raise Conflict("Fichier identique à la version courante", code="SAME_CONTENT")
    if current is None:
        major_no, minor_no = 1, 0
    elif major:
        major_no, minor_no = current.version_major + 1, 0
    else:
        major_no, minor_no = current.version_major, current.version_minor + 1
    version = _add_version(
        db,
        document,
        upload,
        author_id=access.p.id,
        major=major_no,
        minor=minor_no,
        parent=current,
        reason=reason.strip(),
    )
    audit.record(
        db,
        family="business",
        module="documents",
        action="NEW_VERSION",
        resource_type="document",
        resource_id=document.id,
        project_id=document.project_id,
        diff={
            "version": version.label,
            "parent": current.label if current else None,
            "checksum": version.checksum_sha256,
        },
        reason=reason,
    )
    return version


def transition_version(
    db: Session,
    *,
    access: Access,
    version: DocumentVersion,
    target: str,
    comment: str | None = None,
    approvers: list[dict[str, Any]] | None = None,
) -> DocumentVersion:
    document = db.get(Document, version.document_id)
    assert document is not None
    require_read(access, db, document)
    if target not in VERSION_FLOW.get(version.status, set()):
        raise BusinessRuleError(f"Transition {version.status} → {target} interdite", code="INVALID_TRANSITION")
    perm = {
        "in_review": "documents:write",
        "correction": "documents:approve",
        "validated": "documents:approve",
        "published": "documents:publish",
        "obsolete": "documents:publish",
        "archived": "documents:publish",
        "draft": "documents:write",
    }[target]
    access.require(perm)
    previous_status = version.status
    now = utcnow()
    if target == "in_review":
        version.submitted_at = now
        version.locked = True
        document.locked = True
        steps = approvers or []
        for index, step in enumerate(steps):
            db.add(
                DocumentApprovalStep(
                    version_id=version.id,
                    step_order=index + 1,
                    approver_role=step.get("role"),
                    approver_id=uuid.UUID(str(step["user_id"])) if step.get("user_id") else None,
                )
            )
        approver_ids = [str(s["user_id"]) for s in steps if s.get("user_id")]
        emit(
            db,
            "document.approval_requested",
            {"document_id": document.id, "title": document.title, "label": version.label, "approver_ids": approver_ids},
            aggregate_type="document",
            aggregate_id=document.id,
            project_id=document.project_id,
        )
    elif target in ("correction", "validated"):
        pending = (
            db.execute(
                select(DocumentApprovalStep)
                .where(DocumentApprovalStep.version_id == version.id, DocumentApprovalStep.status == "pending")
                .order_by(DocumentApprovalStep.step_order)
            )
            .scalars()
            .all()
        )
        if pending:
            step = pending[0]
            if step.approver_id and step.approver_id != access.p.id and "ADMIN" not in access.p.roles:
                raise Forbidden("Approbation réservée à l'approbateur désigné de l'étape en cours")
            if (
                step.approver_role
                and step.approver_role not in access.p.roles
                and not step.approver_id
                and "ADMIN" not in access.p.roles
            ):
                raise Forbidden("Approbation réservée au rôle désigné")
            step.status = "approved" if target == "validated" else "rejected"
            step.comment = comment
            step.decided_by_id = access.p.id
            step.decided_at = now
            if target == "validated" and len(pending) > 1:
                db.flush()
                audit.record(
                    db,
                    family="business",
                    module="documents",
                    action="APPROVE_STEP",
                    resource_type="document_version",
                    resource_id=version.id,
                    project_id=document.project_id,
                    diff={"step": step.step_order},
                    reason=comment,
                )
                return version  # étape suivante requise
        version.locked = False
        document.locked = False
        if target == "validated":
            version.validated_at = now
    elif target == "published":
        _publish(db, access, document, version)
    elif target == "obsolete":
        version.obsolete_at = now
        version.obsolete_reason = comment
    elif target == "archived":
        version.archived_at = now
    version.status = target
    db.flush()
    audit.record(
        db,
        family="business",
        module="documents",
        action="TRANSITION",
        resource_type="document_version",
        resource_id=version.id,
        project_id=document.project_id,
        diff={"status": [previous_status, target], "version": version.label},
        reason=comment,
        criticality="HIGH" if target == "published" else "LOW",
    )
    return version


def _publish(db: Session, access: Access, document: Document, version: DocumentVersion) -> None:
    """Publication : la version publiée précédente devient obsolète (V4-GED-29) + alertes."""
    now = utcnow()
    previous = db.get(DocumentVersion, document.published_version_id) if document.published_version_id else None
    version.published_at = now
    document.published_version_id = version.id
    document.current_version_id = version.id
    if previous and previous.id != version.id:
        previous.status = "obsolete"
        previous.obsolete_at = now
        previous.obsolete_reason = f"Remplacée par {version.label}"
        previous.superseded_by_id = version.id
        viewers = (
            db.execute(
                select(func.distinct(DocumentAccessLog.user_id)).where(
                    DocumentAccessLog.version_id == previous.id, DocumentAccessLog.user_id.is_not(None)
                )
            )
            .scalars()
            .all()
        )
        if viewers:
            emit(
                db,
                "document.obsolete_alert",
                {
                    "document_id": document.id,
                    "title": document.title,
                    "previous": previous.label,
                    "label": version.label,
                    "user_ids": [str(v) for v in viewers],
                },
                aggregate_type="document",
                aggregate_id=document.id,
                project_id=document.project_id,
            )
    emit(
        db,
        "document.version_published",
        {
            "document_id": document.id,
            "title": document.title,
            "label": version.label,
            "previous": previous.label if previous else None,
        },
        aggregate_type="document",
        aggregate_id=document.id,
        project_id=document.project_id,
    )
    if document.client_visible:
        emit(
            db,
            "document.client_published",
            {"document_id": document.id, "title": document.title, "label": version.label},
            aggregate_type="document",
            aggregate_id=document.id,
            project_id=document.project_id,
        )


def log_access(
    db: Session,
    *,
    document: Document,
    version: DocumentVersion,
    user_id: uuid.UUID,
    action: str,
    application: str | None,
    watermarked: bool,
) -> None:
    db.add(
        DocumentAccessLog(
            document_id=document.id,
            version_id=version.id,
            user_id=user_id,
            action=action,
            application=application,
            watermarked=watermarked,
        )
    )
    if document.confidentiality in ("confidential", "very_sensitive"):
        audit.record(
            db,
            family="business",
            module="documents",
            action=action.upper(),
            resource_type="document_version",
            resource_id=version.id,
            project_id=document.project_id,
            actor_id=user_id,
            criticality="HIGH" if document.confidentiality == "very_sensitive" else "MEDIUM",
        )


def soft_delete(db: Session, *, access: Access, document: Document, reason: str) -> None:
    require_write(access, db, document)
    if len(reason.strip()) < 10:
        raise BusinessRuleError("Justification requise (10 caractères minimum)")
    if document.published_version_id and document.doc_type in ("plan", "contrat", "avenant", "pv", "facture"):
        access.require("documents:publish")
    document.deleted_at = utcnow()
    document.deleted_by_id = access.p.id
    document.deletion_reason = reason
    audit.record(
        db,
        family="business",
        module="documents",
        action="DELETE",
        resource_type="document",
        resource_id=document.id,
        project_id=document.project_id,
        reason=reason,
        criticality="HIGH",
    )


def compare(db: Session, a: DocumentVersion, b: DocumentVersion) -> dict[str, Any]:
    left = (a.extracted_text or "").splitlines()
    right = (b.extracted_text or "").splitlines()
    diff = list(difflib.unified_diff(left, right, fromfile=a.label, tofile=b.label, lineterm="", n=2))
    ratio = difflib.SequenceMatcher(None, "\n".join(left), "\n".join(right)).ratio() if (left or right) else 1.0
    return {
        "from": {
            "id": a.id,
            "label": a.label,
            "checksum": a.checksum_sha256,
            "pages": a.page_count,
            "size": a.size_bytes,
            "created_at": a.created_at,
        },
        "to": {
            "id": b.id,
            "label": b.label,
            "checksum": b.checksum_sha256,
            "pages": b.page_count,
            "size": b.size_bytes,
            "created_at": b.created_at,
        },
        "identical": a.checksum_sha256 == b.checksum_sha256,
        "similarity": round(ratio * 100, 1),
        "text_available": bool(left or right),
        "diff": diff[:3000],
        "added_lines": sum(1 for line in diff if line.startswith("+") and not line.startswith("+++")),
        "removed_lines": sum(1 for line in diff if line.startswith("-") and not line.startswith("---")),
    }


def search(
    db: Session, access: Access, query: str, *, project_id: uuid.UUID | None, limit: int = 30
) -> list[dict[str, Any]]:
    """Recherche plein texte (PostgreSQL, configuration française) dans titres, descriptions et contenus."""
    tsq = func.websearch_to_tsquery("french", query)
    doc_vector = func.to_tsvector(
        "french", func.coalesce(Document.title, "") + " " + func.coalesce(Document.description, "")
    )
    ver_vector = func.to_tsvector(
        "french",
        func.coalesce(DocumentVersion.extracted_text, "") + " " + func.coalesce(DocumentVersion.original_name, ""),
    )
    stmt = (
        select(
            Document,
            DocumentVersion,
            func.ts_headline(
                "french",
                func.coalesce(DocumentVersion.extracted_text, Document.title),
                tsq,
                text("'MaxFragments=2, MaxWords=18, MinWords=6'"),
            ).label("snippet"),
            (func.ts_rank(doc_vector, tsq) * 2 + func.ts_rank(ver_vector, tsq)).label("rank"),
        )
        .join(DocumentVersion, DocumentVersion.id == Document.current_version_id)
        .where(
            Document.deleted_at.is_(None),
            or_(doc_vector.op("@@")(tsq), ver_vector.op("@@")(tsq), Document.code.ilike(f"%{query}%")),
        )
    )
    if project_id:
        stmt = stmt.where(Document.project_id == project_id)
    stmt = stmt.where(Document.confidentiality.in_(access.allowed_confidentialities()))
    rows = db.execute(stmt.order_by(text("rank DESC")).limit(limit * 3)).all()
    out = []
    for document, version, snippet, rank in rows:
        if not can_read(access, db, document):
            continue
        out.append(
            {
                "id": document.id,
                "code": document.code,
                "title": document.title,
                "doc_type": document.doc_type,
                "project_id": document.project_id,
                "version": version.label,
                "status": version.status,
                "snippet": snippet,
                "rank": float(rank or 0),
            }
        )
        if len(out) >= limit:
            break
    return out


def sign_version(
    db: Session,
    *,
    user: Any,
    document: Document,
    version: DocumentVersion,
    otp: dict[str, Any],
    ip: str | None,
    user_agent: str | None,
    role: str | None,
) -> Signature:
    if version.status not in ("validated", "published"):
        raise BusinessRuleError("Seule une version validée ou publiée peut être signée")
    existing = db.execute(
        select(Signature).where(Signature.version_id == version.id, Signature.signer_id == user.id)
    ).scalar_one_or_none()
    if existing:
        raise Conflict("Vous avez déjà signé cette version", code="ALREADY_SIGNED")
    signature = Signature(
        version_id=version.id,
        entity_type="document_version",
        entity_id=version.id,
        signer_id=user.id,
        signer_name=user.full_name,
        signer_email=user.email,
        role=role,
        method="otp",
        signed_hash=version.checksum_sha256,
        otp_channel=otp.get("channel"),
        ip=ip,
        user_agent=(user_agent or "")[:400],
        certificate={
            "otp_verified_at": str(otp.get("verified_at")),
            "document": document.code,
            "version": version.label,
            "sha256": version.checksum_sha256,
        },
    )
    db.add(signature)
    db.flush()
    audit.record(
        db,
        family="business",
        module="signatures",
        action="SIGN",
        resource_type="document_version",
        resource_id=version.id,
        project_id=document.project_id,
        actor_id=user.id,
        diff={"sha256": version.checksum_sha256, "version": version.label},
        criticality="HIGH",
    )
    return signature


def link(
    db: Session,
    *,
    document_id: uuid.UUID,
    entity_type: str,
    entity_id: uuid.UUID,
    version_id: uuid.UUID | None,
    user_id: uuid.UUID,
) -> DocumentLink:
    row = DocumentLink(
        document_id=document_id,
        version_id=version_id,
        entity_type=entity_type,
        entity_id=entity_id,
        created_by_id=user_id,
    )
    db.add(row)
    db.flush()
    return row


def get_document(db: Session, document_id: uuid.UUID) -> Document:
    document = db.get(Document, document_id)
    if document is None or document.deleted_at is not None:
        raise NotFound("Document introuvable")
    return document
