"""Aperçu des documents bureautiques sans téléchargement : Excel (xlsx, xls), OpenDocument (ods, odt, odp), Word
(docx), PowerPoint (pptx), CSV / TSV et texte brut sont convertis en données affichables (tableaux, HTML assaini,
diapositives, texte). Les PDF, images, vidéos et sons sont affichés nativement par le navigateur.

Sécurité : conversion sans état (rien n'est stocké), taille plafonnée, archives ZIP contrôlées (anti « zip bomb »),
XML analysé via defusedxml, HTML produit assaini par nh3, jamais de macro ni de contenu actif exécuté.
"""

from __future__ import annotations

import csv
import html
import io
import zipfile
from pathlib import PurePath
from typing import Any

import nh3
from defusedxml import ElementTree as SafeET

from app.core.errors import AppError

MAX_PREVIEW_BYTES = 25 * 1024 * 1024
MAX_UNZIPPED_BYTES = 200 * 1024 * 1024
MAX_ROWS = 500
MAX_COLS = 60
MAX_SHEETS = 12
MAX_TEXT = 400_000
MAX_SLIDES = 120

TABLE_EXT = {".xlsx", ".xlsm", ".xls", ".ods", ".csv", ".tsv"}
DOC_EXT = {".docx", ".odt"}
SLIDE_EXT = {".pptx", ".odp"}
TEXT_EXT = {".txt", ".md", ".json", ".xml", ".log", ".ifc", ".yaml", ".yml", ".ini", ".sql"}
SUPPORTED = TABLE_EXT | DOC_EXT | SLIDE_EXT | TEXT_EXT

_HTML_TAGS = {
    *("p", "h1", "h2", "h3", "h4", "strong", "em", "u", "s", "ul", "ol", "li", "br", "span", "sup", "sub"),
    *("table", "thead", "tbody", "tr", "th", "td"),
}

NS = {
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


class PreviewError(AppError):
    status_code = 422
    code = "PREVIEW_UNAVAILABLE"


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "isoformat"):
        text = value.isoformat()
        return text[:10] if text.endswith("T00:00:00") else text.replace("T", " ")
    return str(value)[:500]


def _sheet(name: str, rows: list[list[str]], total: int) -> dict[str, Any]:
    while rows and not any(rows[-1]):
        rows.pop()
    width = min(MAX_COLS, max((len(r) for r in rows), default=0))
    # colonnes vides en fin de ligne retirées pour un rendu compact
    while width and not any(len(r) >= width and r[width - 1] for r in rows):
        width -= 1
    return {
        "name": name,
        "columns": width,
        "rows": [(r + [""] * width)[:width] for r in rows],
        "total_rows": total,
        "truncated": total > len(rows),
    }


def _check_zip(data: bytes) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise PreviewError("Fichier endommagé ou format non reconnu") from exc
    if sum(info.file_size for info in archive.infolist()) > MAX_UNZIPPED_BYTES:
        raise PreviewError("Document trop volumineux pour un aperçu : téléchargez-le")
    return archive


# ----------------------------------------------------------------------------------------------- tableurs


def _xlsx(data: bytes) -> dict[str, Any]:
    from openpyxl import load_workbook

    _check_zip(data)
    # Valeurs calculées (cache Excel) ; à défaut (fichier généré par programme, jamais recalculé), la formule.
    values = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    sheets = []
    try:
        for ws, wf in zip(values.worksheets[:MAX_SHEETS], formulas.worksheets[:MAX_SHEETS], strict=False):
            rows: list[list[str]] = []
            total = 0
            for row, raw in zip(ws.iter_rows(values_only=True), wf.iter_rows(values_only=True), strict=False):
                total += 1
                if len(rows) < MAX_ROWS:
                    cells = []
                    for index, value in enumerate(row[:MAX_COLS]):
                        source = raw[index] if index < len(raw) else None
                        if value is None and isinstance(source, str) and source.startswith("="):
                            cells.append(f"ƒ {source}"[:500])
                        else:
                            cells.append(_cell(value))
                    rows.append(cells)
            sheets.append(_sheet(ws.title, rows, total))
    finally:
        values.close()
        formulas.close()
    return {"kind": "table", "sheets": sheets}


def _xls(data: bytes) -> dict[str, Any]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except xlrd.XLRDError as exc:
        raise PreviewError("Classeur Excel illisible") from exc
    sheets = []
    for index in range(min(book.nsheets, MAX_SHEETS)):
        sh = book.sheet_by_index(index)
        rows = []
        for r in range(min(sh.nrows, MAX_ROWS)):
            values = []
            for c in range(min(sh.ncols, MAX_COLS)):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        values.append(_cell(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)))
                        continue
                    except (ValueError, OverflowError):
                        pass
                values.append(_cell(cell.value))
            rows.append(values)
        sheets.append(_sheet(sh.name, rows, sh.nrows))
    return {"kind": "table", "sheets": sheets}


def _ods(data: bytes) -> dict[str, Any]:
    archive = _check_zip(data)
    root = SafeET.fromstring(archive.read("content.xml"))
    sheets = []
    for table in root.iter(f"{{{NS['table']}}}table"):
        if len(sheets) >= MAX_SHEETS:
            break
        rows: list[list[str]] = []
        total = 0
        for row in table.iter(f"{{{NS['table']}}}table-row"):
            repeat = min(int(row.get(f"{{{NS['table']}}}number-rows-repeated", "1")), MAX_ROWS)
            values: list[str] = []
            for cell in row:
                if not cell.tag.endswith("table-cell") and not cell.tag.endswith("covered-table-cell"):
                    continue
                times = min(int(cell.get(f"{{{NS['table']}}}number-columns-repeated", "1")), MAX_COLS)
                text = "\n".join("".join(p.itertext()) for p in cell.iter(f"{{{NS['text']}}}p"))
                values.extend([text[:500]] * times)
                if len(values) >= MAX_COLS:
                    break
            for _ in range(repeat):
                total += 1
                if len(rows) < MAX_ROWS:
                    rows.append(values[:MAX_COLS])
        name = table.get(f"{{{NS['table']}}}name", f"Feuille {len(sheets) + 1}")
        sheets.append(_sheet(name, rows, total))
    return {"kind": "table", "sheets": sheets}


def _decode(data: bytes) -> str:
    for bom, encoding in ((b"\xef\xbb\xbf", "utf-8-sig"), (b"\xff\xfe", "utf-16"), (b"\xfe\xff", "utf-16")):
        if data.startswith(bom):
            return data.decode(encoding, errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    try:
        # exports Excel / logiciels francophones : Windows-1252 avant la détection statistique
        return data.decode("cp1252")
    except UnicodeDecodeError:
        from charset_normalizer import from_bytes

        best = from_bytes(data[:200_000]).best()
        return data.decode(best.encoding if best else "latin-1", errors="replace")


def _csv(data: bytes, ext: str) -> dict[str, Any]:
    text = _decode(data[:MAX_PREVIEW_BYTES])
    sample = text[:20_000]
    if ext == ".tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
        except csv.Error:
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
    rows: list[list[str]] = []
    total = 0
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        total += 1
        if len(rows) < MAX_ROWS:
            rows.append([c[:500] for c in row[:MAX_COLS]])
    return {"kind": "table", "sheets": [_sheet("CSV", rows, total)], "delimiter": delimiter}


# ----------------------------------------------------------------------------------------------- traitement de texte


def _docx(data: bytes) -> dict[str, Any]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_zip(data)
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise PreviewError("Document Word illisible") from exc

    def runs(paragraph: Paragraph) -> str:
        out = []
        for run in paragraph.runs:
            text = html.escape(run.text)
            if not text:
                continue
            if run.bold:
                text = f"<strong>{text}</strong>"
            if run.italic:
                text = f"<em>{text}</em>"
            if run.underline:
                text = f"<u>{text}</u>"
            out.append(text)
        return "".join(out)

    parts: list[str] = []
    list_open: str | None = None
    for block in document.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(block, document)
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            content = runs(paragraph)
            kind = "ol" if "Number" in style else "ul" if ("List" in style or "Bullet" in style) else None
            if kind != list_open and list_open:
                parts.append(f"</{list_open}>")
                list_open = None
            if kind:
                if list_open != kind:
                    parts.append(f"<{kind}>")
                    list_open = kind
                parts.append(f"<li>{content}</li>")
                continue
            if style.startswith("Title"):
                parts.append(f"<h1>{content}</h1>")
            elif style.startswith("Heading"):
                level = min(4, max(2, int("".join(c for c in style if c.isdigit()) or "2") + 1))
                parts.append(f"<h{level}>{content}</h{level}>")
            elif content:
                parts.append(f"<p>{content}</p>")
        elif tag == "tbl":
            if list_open:
                parts.append(f"</{list_open}>")
                list_open = None
            table = Table(block, document)
            rows = []
            for r_index, row in enumerate(table.rows[:MAX_ROWS]):
                cell_tag = "th" if r_index == 0 else "td"
                cells = "".join(f"<{cell_tag}>{html.escape(cell.text)}</{cell_tag}>" for cell in row.cells[:MAX_COLS])
                rows.append(f"<tr>{cells}</tr>")
            parts.append(f"<table><tbody>{''.join(rows)}</tbody></table>")
        if sum(len(p) for p in parts) > MAX_TEXT:
            parts.append("<p><em>… aperçu tronqué : téléchargez le document pour la version complète.</em></p>")
            break
    if list_open:
        parts.append(f"</{list_open}>")
    return {"kind": "html", "html": nh3.clean("".join(parts), tags=_HTML_TAGS, attributes={})}


def _odt(data: bytes) -> dict[str, Any]:
    archive = _check_zip(data)
    root = SafeET.fromstring(archive.read("content.xml"))
    body = root.find(f"{{{NS['office']}}}body")
    parts: list[str] = []
    for element in body.iter() if body is not None else []:
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "h":
            parts.append(f"<h2>{html.escape(''.join(element.itertext()))}</h2>")
        elif tag == "p":
            text = "".join(element.itertext()).strip()
            if text:
                parts.append(f"<p>{html.escape(text)}</p>")
        if sum(len(p) for p in parts) > MAX_TEXT:
            break
    return {"kind": "html", "html": nh3.clean("".join(parts), tags=_HTML_TAGS, attributes={})}


# ----------------------------------------------------------------------------------------------- présentations


def _pptx(data: bytes) -> dict[str, Any]:
    archive = _check_zip(data)
    names = sorted(
        (n for n in archive.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")),
        key=lambda n: int("".join(c for c in PurePath(n).stem if c.isdigit()) or 0),
    )
    slides = []
    for index, name in enumerate(names[:MAX_SLIDES], start=1):
        root = SafeET.fromstring(archive.read(name))
        lines = []
        for paragraph in root.iter(f"{{{NS['a']}}}p"):
            text = "".join(t.text or "" for t in paragraph.iter(f"{{{NS['a']}}}t")).strip()
            if text:
                lines.append(text[:500])
        slides.append({"index": index, "title": lines[0] if lines else f"Diapositive {index}", "lines": lines[1:60]})
    return {"kind": "slides", "slides": slides, "total": len(names)}


def _odp(data: bytes) -> dict[str, Any]:
    archive = _check_zip(data)
    root = SafeET.fromstring(archive.read("content.xml"))
    slides = []
    for index, page in enumerate(root.iter(f"{{{NS['draw']}}}page"), start=1):
        if index > MAX_SLIDES:
            break
        lines = [t for t in ("".join(p.itertext()).strip() for p in page.iter(f"{{{NS['text']}}}p")) if t]
        slides.append({"index": index, "title": lines[0] if lines else f"Diapositive {index}", "lines": lines[1:60]})
    return {"kind": "slides", "slides": slides, "total": len(slides)}


# ----------------------------------------------------------------------------------------------- point d'entrée


def preview(filename: str, data: bytes) -> dict[str, Any]:
    """Convertit un document bureautique en données d'aperçu (aucun stockage)."""
    if not data:
        raise PreviewError("Fichier vide")
    if len(data) > MAX_PREVIEW_BYTES:
        raise PreviewError("Fichier trop volumineux pour un aperçu (25 Mo max) : téléchargez-le", status_code=413)
    ext = PurePath(filename or "").suffix.lower()
    if ext not in SUPPORTED:
        raise PreviewError("Aperçu non disponible pour ce format : téléchargez le fichier", code="PREVIEW_UNSUPPORTED")
    try:
        if ext in (".xlsx", ".xlsm"):
            result = _xlsx(data)
        elif ext == ".xls":
            result = _xls(data)
        elif ext == ".ods":
            result = _ods(data)
        elif ext in (".csv", ".tsv"):
            result = _csv(data, ext)
        elif ext == ".docx":
            result = _docx(data)
        elif ext == ".odt":
            result = _odt(data)
        elif ext == ".pptx":
            result = _pptx(data)
        elif ext == ".odp":
            result = _odp(data)
        else:
            text = _decode(data)
            result = {"kind": "text", "text": text[:MAX_TEXT], "truncated": len(text) > MAX_TEXT}
    except PreviewError:
        raise
    except (KeyError, ValueError, SafeET.ParseError, zipfile.BadZipFile) as exc:
        raise PreviewError("Document illisible ou endommagé : téléchargez-le pour l'ouvrir") from exc
    except Exception as exc:
        raise PreviewError("Aperçu impossible pour ce document : téléchargez-le") from exc
    return {"filename": PurePath(filename).name, "extension": ext.lstrip("."), **result}


__all__ = ["SUPPORTED", "PreviewError", "preview"]
